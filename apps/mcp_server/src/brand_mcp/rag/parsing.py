"""Document parsing into heading-aware blocks: (section path, text)."""

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED = {"md": "md", "markdown": "md", "txt": "txt", "pdf": "pdf", "docx": "docx"}
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class Block:
    section: str  # "Тон голоса > Эмодзи"
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    content_type: str
    blocks: list[Block]


def content_type_for(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower().lstrip(".")
    if suffix not in SUPPORTED:
        raise UnsupportedDocumentError(f"Неподдерживаемый формат: .{suffix or '?'}")
    return SUPPORTED[suffix]


def parse(filename: str, content: bytes) -> ParsedDocument:
    kind = content_type_for(filename)
    stem = PurePath(filename).stem
    if kind in {"md", "txt"}:
        title, blocks = _parse_markdown(content.decode("utf-8", errors="replace"))
    elif kind == "docx":
        title, blocks = _parse_docx(content)
    else:
        title, blocks = _parse_pdf(content)
    blocks = [b for b in blocks if b.text.strip()]
    if not blocks:
        raise UnsupportedDocumentError("В документе не найден текст")
    return ParsedDocument(title=title or stem, content_type=kind, blocks=blocks)


class _SectionBuilder:
    """Accumulates text under the current heading path."""

    def __init__(self) -> None:
        self.path: list[str] = []
        self.title = ""
        self.buffer: list[str] = []
        self.blocks: list[Block] = []

    def heading(self, level: int, text: str) -> None:
        self.flush()
        if level == 1 and not self.title:
            self.title = text
            self.path = []
            return
        depth = max(level - 2, 0)  # H1 is the document title; H2 starts the section path
        self.path = [*self.path[:depth], text]

    def line(self, text: str) -> None:
        self.buffer.append(text)

    def flush(self) -> None:
        text = "\n".join(self.buffer).strip()
        if text:
            self.blocks.append(Block(section=" > ".join(self.path), text=text))
        self.buffer = []


def _parse_markdown(text: str) -> tuple[str, list[Block]]:
    builder = _SectionBuilder()
    for raw in text.splitlines():
        match = _HEADING.match(raw)
        if match:
            builder.heading(len(match.group(1)), match.group(2).strip())
        else:
            builder.line(raw.rstrip())
    builder.flush()
    return builder.title, builder.blocks


def _parse_docx(content: bytes) -> tuple[str, list[Block]]:
    document = DocxDocument(io.BytesIO(content))
    builder = _SectionBuilder()
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        text = paragraph.text.strip()
        if not text:
            builder.line("")
            continue
        if style == "Title":
            builder.heading(1, text)
        elif style.startswith("Heading"):
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 2
            builder.heading(level, text)
        else:
            builder.line(text)
    for table in document.tables:
        for row in table.rows:
            builder.line(" | ".join(cell.text.strip() for cell in row.cells))
    builder.flush()
    return builder.title, builder.blocks


def _parse_pdf(content: bytes) -> tuple[str, list[Block]]:
    reader = PdfReader(io.BytesIO(content))
    blocks = [
        Block(section=f"стр. {number}", text=(page.extract_text() or "").strip())
        for number, page in enumerate(reader.pages, start=1)
    ]
    title = ""
    if reader.metadata is not None and reader.metadata.title:
        title = str(reader.metadata.title)
    return title, blocks
