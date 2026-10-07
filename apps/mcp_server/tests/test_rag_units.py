"""RAG building blocks that need no model: parsing, chunking, rank fusion."""

import io
import uuid

import pytest
from docx import Document as DocxDocument

from brand_mcp.rag.chunking import chunk_blocks
from brand_mcp.rag.parsing import Block, UnsupportedDocumentError, parse
from brand_mcp.rag.search import reciprocal_rank_fusion

MARKDOWN = """# Брендбук Тест

Вступление без раздела.

## Тон голоса

Тёплый и дружеский.

### Эмодзи

Не больше двух эмодзи.

## Цвета

Основной — #4B2E2A.
"""


def test_markdown_sections_follow_heading_hierarchy() -> None:
    parsed = parse("brandbook.md", MARKDOWN.encode())
    assert parsed.title == "Брендбук Тест"
    assert parsed.content_type == "md"
    assert [(b.section, b.text) for b in parsed.blocks] == [
        ("", "Вступление без раздела."),
        ("Тон голоса", "Тёплый и дружеский."),
        ("Тон голоса > Эмодзи", "Не больше двух эмодзи."),
        ("Цвета", "Основной — #4B2E2A."),
    ]


def test_docx_headings_and_tables_are_extracted() -> None:
    doc = DocxDocument()
    doc.add_heading("Брендбук DOCX", level=0)
    doc.add_heading("Шрифты", level=1)
    doc.add_paragraph("Заголовки — Montserrat.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Inter"
    table.rows[0].cells[1].text = "основной текст"
    buffer = io.BytesIO()
    doc.save(buffer)

    parsed = parse("guide.docx", buffer.getvalue())
    assert parsed.title == "Брендбук DOCX"
    text = " ".join(b.text for b in parsed.blocks)
    assert "Montserrat" in text
    assert "Inter | основной текст" in text
    assert parsed.blocks[0].section == "Шрифты"


def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def test_pdf_pages_become_sections() -> None:
    parsed = parse("deck.pdf", _minimal_pdf("Brand colors: navy and lime"))
    assert parsed.content_type == "pdf"
    assert parsed.blocks[0].section == "стр. 1"
    assert "navy and lime" in parsed.blocks[0].text


@pytest.mark.parametrize("filename", ["photo.png", "archive.zip", "noextension"])
def test_unsupported_formats_are_rejected(filename: str) -> None:
    with pytest.raises(UnsupportedDocumentError):
        parse(filename, b"data")


def test_empty_document_is_rejected() -> None:
    with pytest.raises(UnsupportedDocumentError, match="текст"):
        parse("empty.md", b"# Title only\n\n")


def test_chunks_respect_sections_size_and_overlap() -> None:
    sentences = " ".join(f"Предложение номер {i} о кофе." for i in range(60))
    blocks = [Block("Длинный раздел", sentences), Block("Короткий", "Одна строка.")]

    chunks = chunk_blocks(blocks, target_chars=300, overlap_chars=60)

    long_chunks = [c for c in chunks if c.section == "Длинный раздел"]
    assert len(long_chunks) > 3
    assert all(len(c.text) <= 300 + 60 for c in long_chunks)
    # Consecutive chunks share a sentence-aligned overlap.
    last_sentence = long_chunks[0].text.rsplit(". ", 1)[-1]
    assert long_chunks[1].text.startswith(long_chunks[1].text.split(last_sentence)[0])
    assert last_sentence in long_chunks[1].text
    assert "\n\n" not in long_chunks[0].text  # one paragraph stays one paragraph
    assert chunks[-1].section == "Короткий"
    assert chunks[-1].embedding_text() == "Короткий\nОдна строка."


def test_reciprocal_rank_fusion_rewards_agreement() -> None:
    a, b, c = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    scores = reciprocal_rank_fusion([[a, b], [b, c]])
    assert max(scores, key=scores.__getitem__) == b  # ranked by both retrievers
    assert scores[a] == pytest.approx(1 / 61)
    assert scores[b] == pytest.approx(1 / 62 + 1 / 61)
