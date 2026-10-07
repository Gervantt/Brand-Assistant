"""Retrieval over the real demo brand books with the real multilingual embedding model."""

import base64
import io
from pathlib import Path
from typing import Any

import pytest
from conftest import (
    MODEL_CACHE,
    TEST_DATABASE_URL,
    DemoIds,
    connect,
    data,
    text_of,
    token_for,
)
from docx import Document as DocxDocument
from sqlalchemy import delete

from brand_mcp.config import McpSettings
from brand_mcp.deps import Deps
from brand_shared.db.engine import create_engine, create_sessionmaker
from brand_shared.db.models import Document
from brand_shared.permissions import Role

pytestmark = pytest.mark.usefixtures("brandbooks")


async def search(mcp_url: str, ids: DemoIds, role: Role, slug: str, query: str) -> dict[str, Any]:
    async with connect(mcp_url, token_for(ids, role)) as client:
        result = await client.call_tool(
            "search_brandbook", {"client_id": str(ids.clients[slug]), "query": query}
        )
    return data(result)


@pytest.mark.parametrize(
    ("query", "expected"),
    [
        ("Какие фирменные цвета у бренда?", "#4B2E2A"),
        ("Сколько эмодзи можно ставить в посте?", "двух эмодзи"),
        ("Как отвечать на негативные комментарии?", "4 рабочих часов"),
        ("Что такое Bean Card?", "7-й напиток"),
        ("Какие шрифты использовать в заголовках?", "Playfair Display"),
        ("Можно ли упоминать алкоголь?", "Алкоголь"),
    ],
)
async def test_brandbook_questions_find_the_right_fragment(
    mcp_url: str, ids: DemoIds, query: str, expected: str
) -> None:
    result = await search(mcp_url, ids, Role.VIEWER, "bean-there", query)
    assert result["found"] is True
    returned = " ".join(hit["text"] for hit in result["hits"])  # the 5 hits the model sees
    assert expected in returned
    assert 0 < result["confidence"] <= 1


async def test_results_never_leak_across_clients(mcp_url: str, ids: DemoIds) -> None:
    result = await search(mcp_url, ids, Role.MANAGER, "peakform", "Какие фирменные цвета?")
    texts = " ".join(hit["text"] for hit in result["hits"])
    assert "#0B1F3A" in texts
    assert "#4B2E2A" not in texts
    assert all("Bean There" not in hit["document"] for hit in result["hits"])


@pytest.mark.parametrize(
    "query",
    ["Какая столица Франции?", "Сколько стоит iPhone?", "Как настроить Kubernetes кластер?"],
)
async def test_off_topic_questions_are_not_found(mcp_url: str, ids: DemoIds, query: str) -> None:
    result = await search(mcp_url, ids, Role.VIEWER, "bean-there", query)
    assert result["found"] is False


def _docx(text: str) -> str:
    doc = DocxDocument()
    doc.add_heading("Дополнение к брендбуку", level=0)
    doc.add_heading("Секретный ингредиент", level=1)
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode()


async def test_upload_ingest_is_searchable_idempotent_and_permissioned(
    mcp_url: str, ids: DemoIds
) -> None:
    args = {
        "client_id": str(ids.clients["bean-there"]),
        "filename": "addendum.docx",
        "content_base64": _docx("В пряный латте мы добавляем кардамон из Гватемалы."),
    }
    async with connect(mcp_url, token_for(ids, Role.VIEWER)) as client:
        denied = await client.call_tool("ingest_document", args)
    async with connect(mcp_url, token_for(ids, Role.COPYWRITER)) as client:
        first = await client.call_tool("ingest_document", args)
        second = await client.call_tool("ingest_document", args)
        broken = await client.call_tool(
            "ingest_document", {**args, "filename": "photo.png", "content_base64": "AAAA"}
        )

    assert denied.is_error
    assert data(first)["chunks"] >= 1
    assert data(second)["duplicate"] is True
    assert data(second)["document_id"] == data(first)["document_id"]
    assert broken.is_error
    assert "Неподдерживаемый формат" in text_of(broken)

    result = await search(mcp_url, ids, Role.VIEWER, "bean-there", "Какой секретный ингредиент?")
    assert result["found"] is True
    assert "кардамон" in result["hits"][0]["text"]

    # Keep the shared test corpus clean for other retrieval tests.
    engine = create_engine(TEST_DATABASE_URL)
    async with create_sessionmaker(engine)() as session:
        await session.execute(delete(Document).where(Document.id == data(first)["document_id"]))
        await session.commit()
    await engine.dispose()


RERANKER_CACHED = any((Path(MODEL_CACHE)).glob("*jina-reranker-v2*"))


@pytest.mark.skipif(not RERANKER_CACHED, reason="multilingual reranker model not in cache")
@pytest.mark.parametrize(
    ("query", "expected"),
    [
        ("Какой слоган у бренда?", "Свой кофе. Свои люди."),
        ("Что такое Bean Card?", "7-й напиток"),
    ],
)
async def test_reranker_fixes_what_hybrid_search_ranks_low(
    mcp_settings: McpSettings, ids: DemoIds, query: str, expected: str
) -> None:
    settings = mcp_settings.model_copy(update={"reranker_enabled": True})
    deps = Deps.create(settings)
    async with deps.sessionmaker() as session:
        hits = await deps.retriever.search(session, ids.clients["bean-there"], query, 3)
    await deps.aclose()
    assert expected in hits[0].text
    assert hits[0].score >= settings.min_rerank_score
